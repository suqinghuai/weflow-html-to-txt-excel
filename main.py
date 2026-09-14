import os
import re
import json
import glob
import html
import sys
from datetime import datetime
from html.parser import HTMLParser


def get_app_dir():
    """返回程序所在目录，而不是当前工作目录或 PyInstaller 临时目录。"""
    if getattr(sys, 'frozen', False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


APP_DIR = get_app_dir()


class MessageContentParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text_parts = []
        self.quoted_sender = None
        self.quoted_text = None
        self.link_title = None
        self.link_url = None
        self.has_message_text = False
        self.has_quoted_message = False
        self.has_link_card = False
        self._in_message_text = False
        self._in_quoted_sender = False
        self._in_quoted_text = False
        self._in_link_card = False
        self._current_tag_stack = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        class_attr = attrs_dict.get('class', '')

        if 'message-text' in class_attr:
            self._in_message_text = True
            self.has_message_text = True
        elif 'quoted-message' in class_attr:
            self.has_quoted_message = True
        elif 'quoted-sender' in class_attr:
            self._in_quoted_sender = True
        elif 'quoted-text' in class_attr:
            self._in_quoted_text = True
        elif 'message-link-card' in class_attr:
            self._in_link_card = True
            self.has_link_card = True
            self.link_url = attrs_dict.get('href', '')

    def handle_endtag(self, tag):
        if self._in_message_text and tag == 'div':
            self._in_message_text = False
        if self._in_quoted_sender and tag == 'div':
            self._in_quoted_sender = False
        if self._in_quoted_text and tag == 'div':
            self._in_quoted_text = False
        if self._in_link_card and tag == 'a':
            self._in_link_card = False

    def handle_data(self, data):
        if self._in_message_text:
            self.text_parts.append(data)
        elif self._in_quoted_sender:
            self.quoted_sender = data
        elif self._in_quoted_text:
            self.quoted_text = data
        elif self._in_link_card:
            self.link_title = data

    def get_text(self):
        return ''.join(self.text_parts)


def parse_message_body(body_html):
    parser = MessageContentParser()
    parser.feed(body_html)
    return parser


def classify_message(parser_result, raw_text):
    if parser_result.has_quoted_message:
        return '引用消息'

    if parser_result.has_message_text and raw_text:
        if re.search(r'\[语音通话\]|\[视频通话\]', raw_text):
            return '通话消息'
        if '[表情包]' in raw_text:
            return '动画表情'
        if '撤回了一条消息' in raw_text:
            return '系统消息'
        if '[转账]' in raw_text or '拍了拍' in raw_text:
            return '其他消息'
        if parser_result.has_link_card:
            return '其他消息'
        return '文本消息'

    if parser_result.has_link_card:
        return '其他消息'

    return '图片消息'


def build_txt_content(raw_text, parser_result, msg_type):
    parts = []

    if parser_result.has_quoted_message:
        quote_suffix = ''
        if parser_result.quoted_sender and parser_result.quoted_text:
            quote_suffix = f'[引用 {parser_result.quoted_sender}：{parser_result.quoted_text}]'
        elif parser_result.quoted_text:
            quote_suffix = f'[引用 {parser_result.quoted_text}]'
        if raw_text:
            parts.append(f'{raw_text}{quote_suffix}')
        else:
            parts.append(quote_suffix)
    elif parser_result.has_link_card:
        link_title = parser_result.link_title or ''
        if raw_text and raw_text != link_title:
            parts.append(raw_text)
        link_line = f'[链接] {link_title}' if link_title else '[链接]'
        parts.append(link_line)
        if parser_result.link_url:
            parts.append(html.unescape(parser_result.link_url))
    elif raw_text:
        parts.append(raw_text)
    else:
        if msg_type == '图片消息':
            parts.append('[图片]')
        elif msg_type == '视频消息':
            parts.append('[视频]')

    return '\n'.join(parts)


def build_xlsx_content(raw_text, parser_result):
    if parser_result.has_quoted_message:
        quote_suffix = ''
        if parser_result.quoted_sender and parser_result.quoted_text:
            quote_suffix = f'[引用 {parser_result.quoted_sender}：{parser_result.quoted_text}]'
        elif parser_result.quoted_text:
            quote_suffix = f'[引用 {parser_result.quoted_text}]'
        return f'{raw_text}{quote_suffix}' if raw_text else quote_suffix

    if parser_result.has_link_card:
        if raw_text:
            return raw_text
        return parser_result.link_title if parser_result.link_title else '[链接]'

    if raw_text:
        return raw_text

    msg_type = classify_message(parser_result, raw_text)
    if msg_type == '图片消息':
        return '[图片]'
    if msg_type == '视频消息':
        return '[视频]'

    return ''


def extract_avatar_info(avatar_html):
    alt_match = re.search(r'alt="([^"]*)"', avatar_html)
    src_match = re.search(r'src="([^"]*)"', avatar_html)
    alt = alt_match.group(1) if alt_match else ''
    src = src_match.group(1) if src_match else ''
    wxid_match = re.search(r'avatars/([^/]+)\.', src)
    wxid = wxid_match.group(1) if wxid_match else ''
    return alt, wxid


def extract_time_from_body(body_html):
    time_match = re.search(r'class="message-time">([^<]+)<', body_html)
    return time_match.group(1) if time_match else ''


def parse_html_file(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    title_match = re.search(r'<h1 class="title">([^<]+)</h1>', content)
    chat_title = title_match.group(1) if title_match else ''

    meta_match = re.search(r'<span>(\d+) 条消息</span>', content)
    msg_count = meta_match.group(1) if meta_match else ''

    type_match = re.search(r'<span>(私聊|群聊)</span>', content)
    chat_type = type_match.group(1) if type_match else '私聊'

    data_match = re.search(r'window\.WEFLOW_DATA\s*=\s*(\[.*?\]);', content, re.DOTALL)
    if not data_match:
        print(f'  [跳过] 未找到 WEFLOW_DATA: {os.path.basename(filepath)}')
        return None

    messages_data = json.loads(data_match.group(1))

    other_name = ''
    other_wxid = ''
    my_name = ''
    my_wxid = ''
    my_names_set = set()

    messages = []
    for item in messages_data:
        time_str = extract_time_from_body(item['b'])
        sender_alt, sender_wxid = extract_avatar_info(item['a'])
        is_me = item.get('s', 0) == 1

        if is_me:
            if not my_name and sender_alt:
                my_name = sender_alt
            if not my_wxid and sender_wxid:
                my_wxid = sender_wxid
            if sender_alt:
                my_names_set.add(sender_alt)
        else:
            if not other_name and sender_alt:
                other_name = sender_alt
            if not other_wxid and sender_wxid:
                other_wxid = sender_wxid

        parser_result = parse_message_body(item['b'])
        raw_text = parser_result.get_text()

        if parser_result.has_quoted_message and parser_result.quoted_sender:
            if is_me and parser_result.quoted_sender != other_name:
                my_names_set.add(parser_result.quoted_sender)
            elif not is_me and parser_result.quoted_sender != other_name:
                my_names_set.add(parser_result.quoted_sender)

        msg_type = classify_message(parser_result, raw_text)
        txt_content = build_txt_content(raw_text, parser_result, msg_type)
        xlsx_content = build_xlsx_content(raw_text, parser_result)

        sender_label = '我' if is_me else other_name

        messages.append({
            'index': item['i'],
            'time': time_str,
            'sender': sender_label,
            'is_me': is_me,
            'msg_type': msg_type,
            'txt_content': txt_content,
            'xlsx_content': xlsx_content,
        })

    account_nickname = ''
    if my_names_set:
        for name in my_names_set:
            if name != my_name:
                account_nickname = name
                break
        if not account_nickname:
            account_nickname = my_name
    else:
        account_nickname = my_name

    export_time = datetime.now().strftime('%Y-%m-%d %H:%M:%S')

    return {
        'chat_title': chat_title,
        'chat_type': chat_type,
        'msg_count': msg_count,
        'other_name': other_name,
        'other_wxid': other_wxid,
        'my_name': my_name,
        'my_wxid': my_wxid,
        'account_nickname': account_nickname,
        'export_time': export_time,
        'messages': messages,
    }


def write_txt(filepath, data):
    with open(filepath, 'w', encoding='utf-8') as f:
        for msg in data['messages']:
            f.write(f"{msg['time']} '{msg['sender']}'\n")
            if msg['txt_content']:
                f.write(f"{msg['txt_content']}\n")
            f.write('\n')


ILLEGAL_XML_CHARS_RE = re.compile(
    r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x84\x86-\x9f]'
)


def sanitize_xlsx_value(value):
    if isinstance(value, str):
        return ILLEGAL_XML_CHARS_RE.sub('', value)
    return value


def write_xlsx(filepath, data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = '聊天记录'

    header_font = Font(bold=True, size=11)
    center_align = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin'),
    )
    header_fill = PatternFill(start_color='D9E1F2', end_color='D9E1F2', fill_type='solid')

    ws.merge_cells('A1:H1')
    ws['A1'] = '会话信息'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = center_align

    ws['A2'] = '微信ID'
    ws['B2'] = sanitize_xlsx_value(data['other_wxid'])
    ws['D2'] = '昵称'
    ws['E2'] = sanitize_xlsx_value(data['account_nickname'])
    for cell in [ws['A2'], ws['D2']]:
        cell.font = header_font

    ws['A3'] = '导出工具'
    ws['B3'] = 'WeFlow'
    ws['C3'] = '导出版本'
    ws['D3'] = '0.0.2'
    ws['E3'] = '平台'
    ws['F3'] = 'wechat'
    ws['G3'] = '导出时间'
    ws['H3'] = data['export_time']
    for cell in [ws['A3'], ws['C3'], ws['E3'], ws['G3']]:
        cell.font = header_font

    column_headers = ['序号', '时间', '发送者身份', '消息类型', '内容']
    for col_idx, header in enumerate(column_headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.alignment = center_align
        cell.fill = header_fill
        cell.border = thin_border

    for row_idx, msg in enumerate(data['messages'], 5):
        values = [msg['index'], msg['time'], msg['sender'], msg['msg_type'], msg['xlsx_content']]
        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=sanitize_xlsx_value(value))
            cell.border = thin_border
            if col_idx == 1:
                cell.alignment = center_align

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 22
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 60

    wb.save(filepath)


def main():
    html_files = [
        filepath for filepath in glob.glob(os.path.join(APP_DIR, '*'))
        if os.path.isfile(filepath) and os.path.splitext(filepath)[1].lower() == '.html'
    ]

    if not html_files:
        print('未找到任何 HTML 文件')
        return

    print(f'找到 {len(html_files)} 个 HTML 文件\n')

    for html_path in html_files:
        basename = os.path.splitext(os.path.basename(html_path))[0]
        print(f'处理: {basename}.html')

        data = parse_html_file(html_path)
        if not data:
            continue

        txt_path = os.path.join(APP_DIR, f'{basename}.txt')
        xlsx_path = os.path.join(APP_DIR, f'{basename}.xlsx')

        write_txt(txt_path, data)
        print(f'  -> {basename}.txt ({len(data["messages"])} 条消息)')

        write_xlsx(xlsx_path, data)
        print(f'  -> {basename}.xlsx ({len(data["messages"])} 条消息)')

    print('\n转换完成！')


def wait_for_exit():
    print('\n按任意键退出程序...', end='', flush=True)
    if os.name == 'nt':
        import msvcrt
        msvcrt.getch()
    else:
        input()


if __name__ == '__main__':
    try:
        main()
    except Exception as error:
        print(f'\n程序运行失败: {error}')
    finally:
        wait_for_exit()